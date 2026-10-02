#!/usr/bin/env python3
"""Mathematics IX — Part II allocation (founder ruling, 2026-10-02).

WHAT WAS DECIDED
  * Part I (ch 1-8) keeps the period counts its certified libraries were authored at:
    10, 12, 17, 15, 15, 15, 10, 12 = 106. Re-running the allocation over all 14 real
    effort indices would add 1-2 periods to every Part I chapter (116 in all) and put each
    one above its authored top canonical, so the serve engine would surrender the extra
    periods. The change is too small to be worth re-authoring.
  * Part II (ch 9-14, NCERT Ganita Manjari Part II, released Aug 2026) takes the counts
    the full re-run gives it: 16, 16, 14, 16, 16, 16 = 94. The 13-period placeholders were
    only ever a stand-in.
  * 106 + 94 = 200 of the 210-period year. The remaining 10 are RESERVED FOR PRACTICE —
    a Year Plan row, never a chapter (genon/master_plan.py RESERVE_PREFIX).

HOW IT IS ENCODED
  The app spreads a teacher's own annual budget by each row's weight (largest remainder),
  so for this one combo the weights ARE the ruled period counts: at the calibrated 210 the
  suggestions equal the authored tops exactly (no surrender), and any other budget scales
  the same plan proportionally, reserve included. So for mathematics·IX the Weight column
  holds calibrated periods, not effort indices — every other combo is untouched and keeps
  weight == effort index. The effort indices stay in the mapping files (Part I 79.5,
  Part II 65.5) and are repeated in a comment on each weight cell.

    python3 genon/out/stage_prep_mathematics_secondary/apply_ix_part2_allocation.py           # dry run
    python3 genon/out/stage_prep_mathematics_secondary/apply_ix_part2_allocation.py --apply
    ... --workbook /path/to/copy.xlsx --apply      # rehearse on a copy (no backups taken)
    ... --apply --lock-checked                     # Excel confirmed closed, stale ~$ lock left behind

AFTERWARDS, and it is not optional (the runbook pair):
    python3 genon/master_plan.py && python3 genon/variant_plans.py
master_plan.py rebuilds every row from the workbook and WIPES canonical_plan;
variant_plans.py must re-annotate immediately or no row is trustworthy.
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[3]
NORMS = ROOT / "data/cloud/content/allocation_norms"
BACKUP = ROOT / "backup/allocation_workbook"

SUBJECT, CLS = "Mathematics", "IX"
RESERVE_TITLE = "Reserved for practice"          # must start with master_plan.RESERVE_PREFIX
RESERVE_PERIODS = 10

# (current effort index, ruled periods) — Part I indices from the 2026-08 mappings,
# Part II from the 2026-10-02 mappings.
RULED = {
    1: (7.5, 10), 2: (9, 12), 3: (13, 17), 4: (11, 15),
    5: (11.5, 15), 6: (11, 15), 7: (7.5, 10), 8: (9, 12),
    9: (11.5, 16), 10: (11, 16), 11: (9.5, 14), 12: (11.5, 16), 13: (11, 16), 14: (11, 16),
}
TITLES = {
    9: "Propositions and their Converses",
    10: "How Quantities Combine: Understanding Data",
    11: "The World of Algorithms",
    12: "Quadrilaterals",
    13: "Two Variables, One Line",
    14: "Math of Space: Surface Area and Volume",
}
PLACEHOLDER_WEIGHT = 9.9375
AUTHOR = "Aruvi allocation ruling 2026-10-02"

args = sys.argv[1:]
apply = "--apply" in args
wb_path = Path(args[args.index("--workbook") + 1]) if "--workbook" in args else NORMS / "ncf_chapterwise_period_allocation.xlsx"
real = wb_path.resolve() == (NORMS / "ncf_chapterwise_period_allocation.xlsx").resolve()

lock = wb_path.parent / f"~${wb_path.name}"
if real and apply and lock.exists():
    if "--lock-checked" not in args:
        raise SystemExit(f"ABORT: {lock.name} exists — the workbook may be open in Excel. Close it "
                         "and run again. If Excel is closed and the lock is stale (macOS sometimes "
                         "leaves one behind), re-run with --lock-checked.")
    print(f"NOTE: {lock.name} is present but Excel was confirmed closed — treating it as stale.")

# data_only=False so the Summary sheet's formulas survive the round-trip.
wb = openpyxl.load_workbook(wb_path, data_only=False)
ws = wb["Chapters"]
rows = {r[2].value: r for r in ws.iter_rows(min_row=2) if r[0].value == SUBJECT and r[1].value == CLS}

budget = next(r[2] for r in wb["budget"].iter_rows(min_row=2, values_only=True)
              if r[0] == SUBJECT and r[1] == CLS)
if sum(p for _, p in RULED.values()) + RESERVE_PERIODS != budget:
    raise SystemExit(f"ABORT: ruled periods + reserve != the {budget}-period budget")

already = (15 in rows and str(rows[15][3].value) == RESERVE_TITLE and 16 not in rows
           and all(rows[ch][4].value == p for ch, (_, p) in RULED.items()))
if already:
    print("Already applied — nothing to do.")
    sys.exit(0)

# Expected starting state: Part I at its effort indices, 9-16 placeholders.
if sorted(rows) != list(range(1, 17)):
    raise SystemExit(f"ABORT: expected maths IX chapters 1-16 in the workbook, found {sorted(rows)}")
for ch in range(1, 9):
    if rows[ch][4].value != RULED[ch][0]:
        raise SystemExit(f"ABORT: ch{ch} weight is {rows[ch][4].value!r}, expected effort index "
                         f"{RULED[ch][0]} — the workbook has moved since this script was written.")
for ch in range(9, 17):
    if "Placeholder" not in str(rows[ch][3].value) or rows[ch][4].value != PLACEHOLDER_WEIGHT:
        raise SystemExit(f"ABORT: ch{ch} is not the expected placeholder row: "
                         f"{rows[ch][3].value!r} / {rows[ch][4].value!r}")

print(f"{'ch':>3}  {'weight now':>10} -> {'new':>4}   title")
for ch in range(1, 15):
    old_w, (ei, p) = rows[ch][4].value, RULED[ch]
    title = TITLES.get(ch, rows[ch][3].value)
    print(f"{ch:>3}  {old_w!s:>10} -> {p:>4}   {title}" + ("   (title set)" if ch in TITLES else ""))
print(f" 15  {rows[15][4].value!s:>10} -> {RESERVE_PERIODS:>4}   {RESERVE_TITLE}   (reserve row)")
print(f" 16  {rows[16][4].value!s:>10} ->  del   row removed ({rows[16][3].value})")
print(f"Total: {sum(p for _, p in RULED.values())} chapter periods + {RESERVE_PERIODS} reserve = {budget}")

if not apply:
    print("\nDRY RUN — nothing written. Re-run with --apply.")
    sys.exit(0)

if real:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    BACKUP.mkdir(parents=True, exist_ok=True)
    for src in (wb_path, NORMS / "master_plan.json", NORMS / "canonical_period_pins.json"):
        dst = BACKUP / f"{src.stem}_{ts}{src.suffix}"
        shutil.copy2(src, dst)
        print(f"backup: {dst.relative_to(ROOT)}")

for ch in range(1, 15):
    ei, p = RULED[ch]
    cell = rows[ch][4]
    cell.value = p
    cell.comment = Comment(
        f"Calibrated periods, not an effort index (founder ruling 2026-10-02). "
        f"Effort index {ei:g}. Part I keeps its authored counts (106); Part II takes the "
        f"full re-run counts (94); 10 periods are reserved for practice.", AUTHOR)
    if ch in TITLES:
        rows[ch][3].value = TITLES[ch]
rows[15][3].value = RESERVE_TITLE
rows[15][4].value = RESERVE_PERIODS
rows[15][4].comment = Comment(
    "Budgeted practice time: 10 of the 210 periods. Not a chapter — no summary, mapping or "
    "canonicals (genon/master_plan.py RESERVE_PREFIX). Founder ruling 2026-10-02.", AUTHOR)
ws.delete_rows(rows[16][0].row, 1)

wb.save(wb_path)
print(f"\nwrote {wb_path}")
print("NEXT: python3 genon/master_plan.py && python3 genon/variant_plans.py")
