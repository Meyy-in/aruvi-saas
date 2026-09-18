"""A paid subject with NO classes must still appear in the data-rights export (app. 04 E2).

Since 5d·11 a subject she has paid for survives losing its last class, so `grades: []` is a
real record. The export walked only the grades, so the subject vanished from her teaching
profile table — and when it was her only subject the document said "No teaching profile on
record", a false statement about a subscriber's record. Stdlib + python-docx only.
"""
import io, sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from docx import Document
from aruvi_core.export_data_rights_docx import build_export_docx


def _text(payload):
    d = Document(io.BytesIO(build_export_docx(payload)))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text for c in r.cells))
    return "\n".join(parts)


def test_only_subject_with_no_classes_is_listed_not_denied():
    out = _text({"account": {}, "profile": {"subjects": [{"name": "English", "grades": []}]},
                 "years": []})
    assert "No teaching profile on record" not in out
    assert "English | — | —" in out


def test_classless_subject_listed_beside_a_taught_one():
    out = _text({"account": {}, "profile": {"subjects": [
        {"name": "Science", "grades": [{"grade": "IX", "sections": [{"tag": "9A"}]}]},
        {"name": "English", "grades": []}]}, "years": []})
    assert "Science | 9 | 9A" in out or "Science | IX | 9A" in out
    assert "English | — | —" in out


def test_empty_profile_still_says_so():
    out = _text({"account": {}, "profile": {"subjects": []}, "years": []})
    assert "No teaching profile on record" in out


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f(); print("ok", n)
